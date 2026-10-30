"""Ghi lại mọi thao tác kéo vào file Excel (và Google Sheets nếu bật)."""
import os
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

import excel_view

HEADERS = ["Thời gian", "Link video / Kênh", "Hạng mục", "Chi tiết",
           "Trạng thái", "Lỗi", "Ảnh chụp lỗi", "Tình trạng gỡ", "Kênh",
           "Kênh mạnh (dùng để kéo)", "Kênh cần kéo", "Video đã gắn", "File kênh",
           "Dữ liệu gắn (JSON)", "Thời gian gỡ"]
WIDTHS = [19, 48, 22, 45, 18, 45, 30, 16, 24, 34, 34, 40, 16, 60, 19]
COL = {h: i + 1 for i, h in enumerate(HEADERS)}

GREEN = PatternFill("solid", start_color="C6EFCE")
RED = PatternFill("solid", start_color="FFC7CE")
YELLOW = PatternFill("solid", start_color="FFEB9C")

STATUS_DONE = "✅ DONE"
STATUS_ERROR = "❌ LỖI"
STATUS_CHECK = "⚠️ CẦN KIỂM TRA"
REMOVE_NOT = "Chưa gỡ"
REMOVE_DONE = "✅ Đã gỡ"
REMOVE_ERROR = "❌ Lỗi khi gỡ"
REMOVE_MANUAL = "⚠️ Gỡ thủ công"
REMOVE_SKIP = "⏭ Bỏ qua – không thuộc kênh mạnh"   # mục không phải của kênh mạnh này -> không bao giờ đụng

FILLS = {STATUS_DONE: GREEN, STATUS_ERROR: RED, STATUS_CHECK: YELLOW,
         REMOVE_DONE: GREEN, REMOVE_ERROR: RED, REMOVE_MANUAL: YELLOW}


class Tracker:
    def __init__(self, cfg: dict):
        self.path = cfg["tracker"]["excel_file"]
        self.session = {}  # số thao tác theo trạng thái trong LẦN CHẠY NÀY
        self.job = cfg.get("job_name") or cfg.get("weak_channel", {}).get("name", "")
        s, w = cfg.get("strong_channel", {}), cfg.get("weak_channel", {})
        self.strong = f"{s.get('name', '')} ({s.get('channel_id', '')})".strip()
        self.weak = f"{w.get('name', '')} ({w.get('channel_id', '')})".strip()
        self.job_file = os.path.basename(cfg.get("_path", "") or "")
        self.gs_cfg = cfg["tracker"].get("google_sheet", {}) or {}
        self.saved_to = self.path
        self._warned = False
        excel_view.merge_pending(self.path)          # lần trước file đang mở trong WPS -> gộp phần ghi tạm
        src = excel_view.latest(self.path)
        if os.path.exists(src):
            self.wb = load_workbook(src)
            self.ws = excel_view.log_sheet(self.wb)
            changed = False
            for name in HEADERS:  # file cũ: bổ sung các cột mới
                c = self.ws.cell(1, COL[name])
                if c.value != name:
                    c.value = name
                    c.font = Font(name="Arial", bold=True, color="FFFFFF")
                    c.fill = PatternFill("solid", start_color="305496")
                    self.ws.column_dimensions[c.column_letter].width = WIDTHS[COL[name] - 1]
                    changed = True
            if changed or excel_view.VIEW_NAME not in self.wb.sheetnames:   # file cũ -> thêm trang Kết quả gọn
                self.save()
        else:
            self.wb = Workbook()
            self.ws = self.wb.active
            self.ws.title = excel_view.DATA_NAME
            self.ws.append(HEADERS)
            for c in self.ws[1]:
                c.font = Font(name="Arial", bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", start_color="305496")
                c.alignment = Alignment(horizontal="center")
            for i, w in enumerate(WIDTHS, 1):
                self.ws.column_dimensions[self.ws.cell(1, i).column_letter].width = w
            self.ws.freeze_panes = "A2"
            self.save()

    # ---------- ghi ----------
    def log(self, target, task, detail, status, error="", screenshot="", attached="", data=""):
        """attached: mã các video đã gắn vào video này; data: JSON chi tiết từng vị trí (để đối soát / gỡ)."""
        import json
        import re
        if data and not attached:
            try:
                attached = ", ".join(dict.fromkeys(x["id"] for x in json.loads(data) if x.get("id")))
            except Exception:
                pass
        self.session[status] = self.session.get(status, 0) + 1
        removal = REMOVE_NOT if status != STATUS_ERROR else ""
        if not attached and detail:
            attached = ", ".join(dict.fromkeys(re.findall(r"\(([\w-]{11})\)", str(detail))))
        self.ws.append([datetime.now().strftime("%Y-%m-%d %H:%M:%S"), target, task,
                        detail, status, error, screenshot, removal, self.job,
                        self.strong, self.weak, attached, self.job_file, data, ""])
        row = self.ws.max_row
        for c in self.ws[row]:
            c.font = Font(name="Arial")
        self._paint(row, COL["Trạng thái"], status)
        self._paint(row, COL["Tình trạng gỡ"], removal)
        self.save()
        print(f"  {status:<16} {task:<28} {target}" +
              (f"\n      ↳ {detail}" if detail and status != STATUS_ERROR else "") +
              (f"\n      → {error}" if error else ""), flush=True)

    def set_removal(self, row, value, error=""):
        self.ws.cell(row=row, column=COL["Tình trạng gỡ"], value=value)
        self.ws.cell(row=row, column=COL["Thời gian gỡ"], value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
        if error:
            old = self.ws.cell(row=row, column=COL["Lỗi"]).value or ""
            self.ws.cell(row=row, column=COL["Lỗi"], value=(old + " | GỠ: " + error).strip(" |"))
        self._paint(row, COL["Tình trạng gỡ"], value)
        self.save()

    # ---------- đọc ----------
    def rows_to_revert(self, only_rows=None):
        """Các dòng đã kéo thành công nhưng chưa gỡ (chỉ trong `only_rows` nếu có)."""
        out = []
        for r in range(2, self.ws.max_row + 1):
            if only_rows is not None and r not in only_rows:
                continue
            status = self.ws.cell(r, COL["Trạng thái"]).value
            removal = self.ws.cell(r, COL["Tình trạng gỡ"]).value
            detail = str(self.ws.cell(r, COL["Chi tiết"]).value or "")
            if status in (STATUS_DONE, STATUS_CHECK) and removal in (REMOVE_NOT, REMOVE_ERROR, REMOVE_MANUAL) \
                    and "đã có sẵn" not in detail and "đã đúng sẵn trên kênh" not in detail:  # video có sẵn từ trước, tool không gắn gì -> không gỡ
                out.append({"row": r,
                            "target": self.ws.cell(r, COL["Link video / Kênh"]).value,
                            "task": self.ws.cell(r, COL["Hạng mục"]).value,
                            "detail": detail,
                            "attached": self.ws.cell(r, COL["Video đã gắn"]).value or "",
                            "file": self.ws.cell(r, COL["File kênh"]).value or "",
                            "job": self.ws.cell(r, COL["Kênh"]).value or "",
                            "strong": str(self.ws.cell(r, COL["Kênh mạnh (dùng để kéo)"]).value or ""),
                            "weak": str(self.ws.cell(r, COL["Kênh cần kéo"]).value or ""),
                            "data":self.ws.cell(r, COL["Dữ liệu gắn (JSON)"]).value or ""})
        return out

    def active_rows(self):
        """Các thao tác đã làm (DONE / CẦN KIỂM TRA) và CHƯA bị gỡ."""
        out = []
        for r in range(2, self.ws.max_row + 1):
            status = self.ws.cell(r, COL["Trạng thái"]).value
            removal = self.ws.cell(r, COL["Tình trạng gỡ"]).value
            if status in (STATUS_DONE, STATUS_CHECK) and removal in (REMOVE_NOT, REMOVE_ERROR, None, ""):
                out.append({"target": self.ws.cell(r, COL["Link video / Kênh"]).value,
                            "task": self.ws.cell(r, COL["Hạng mục"]).value,
                            "status": status,
                            "attached": str(self.ws.cell(r, COL["Video đã gắn"]).value or ""),
                            "note": (self.ws.cell(r, COL["Chi tiết"]).value or "") + " " +
                                    (self.ws.cell(r, COL["Lỗi"]).value or "")})
        return out

    def summary(self):
        counts = {}
        for r in range(2, self.ws.max_row + 1):
            s = self.ws.cell(r, COL["Trạng thái"]).value
            counts[s] = counts.get(s, 0) + 1
        return counts

    # ---------- nội bộ ----------
    def _paint(self, row, col, value):
        fill = FILLS.get(value)
        if fill:
            self.ws.cell(row=row, column=col).fill = fill

    def save(self):
        try:
            excel_view.tidy_data_sheet(self.ws)
            excel_view.build_view(self.wb)   # trang "Kết quả" gọn, dễ đọc – dựng lại từ dữ liệu gốc
        except Exception as e:                # trang xem lỗi cũng không được làm mất nhật ký
            print(f"⚠️ Không dựng được trang Kết quả: {e}", flush=True)
        try:
            self.wb.save(self.path)
            self.saved_to = self.path
            excel_view.drop_pending(self.path)
        except PermissionError:
            # file đang mở trong Excel / WPS: KHÔNG để mất kết quả -> ghi đầy đủ sang file phụ, đóng WPS là tự gộp
            tmp = excel_view.pending_path(self.path)
            try:
                self.wb.save(tmp)
                self.saved_to = tmp
            except Exception as e:
                print(f"❌ Không lưu được nhật ký (cả file phụ): {e}", flush=True)
                return
            if not self._warned:
                self._warned = True
                print(f"⚠️ File {os.path.basename(self.path)} đang MỞ trong Excel / WPS nên tool chưa ghi đè được. "
                      f"Kết quả vẫn được lưu đủ vào file phụ {os.path.basename(tmp)} – ĐÓNG file Excel / WPS, "
                      f"tool tự gộp vào file chính.", flush=True)

    def sync_google_sheet(self):
        if not self.gs_cfg.get("enabled"):
            return
        try:
            import gspread
            gc = gspread.service_account(filename=self.gs_cfg["credentials_json"])
            ws = gc.open_by_url(self.gs_cfg["sheet_url"]).sheet1
            data = [[("" if v is None else str(v)) for v in row]
                    for row in self.ws.iter_rows(values_only=True)]
            ws.clear()
            ws.update(range_name="A1", values=data)
            print("☁️  Đã đồng bộ lên Google Sheets.")
        except Exception as e:  # không để lỗi đồng bộ làm hỏng cả phiên chạy
            print(f"⚠️  Không đồng bộ được Google Sheets: {e}")
