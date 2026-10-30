"""
Lịch sử của module Phân tích & tương tác video (SQLite: phan_tich.db cạnh tool).
Mỗi video đã phân tích = 1 dòng; mỗi lượt chạy = 1 run.
"""
import csv
import json
import os
import sqlite3
import threading
from datetime import datetime

BASE = os.path.dirname(os.path.abspath(__file__))
import paths   # noqa: E402
DB_PATH = os.path.join(paths.DATA, "phan_tich.db")

COLUMNS = ("STT", "Video", "Chủ đề", "Điểm", "Đề xuất", "Trạng thái")

SCHEMA = """
CREATE TABLE IF NOT EXISTS runs (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    started_at TEXT, finished_at TEXT, account TEXT, target TEXT, keywords TEXT, settings TEXT, note TEXT
);
CREATE TABLE IF NOT EXISTS videos (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    run_id INTEGER,
    video_id TEXT, url TEXT, title TEXT, channel TEXT, channel_url TEXT,
    duration INTEGER, views INTEGER, likes INTEGER,
    lang TEXT, relevance REAL, topic TEXT,
    window_from INTEGER, window_to INTEGER, transcript_source TEXT, transcript_chars INTEGER,
    score INTEGER, category TEXT, recommendation TEXT, reasons TEXT, dislike_basis TEXT,
    comments TEXT, comment_lang TEXT, ai_provider TEXT,
    status TEXT, action TEXT, action_at TEXT, error TEXT,
    found_at TEXT, analyzed_at TEXT
);
CREATE INDEX IF NOT EXISTS idx_videos_vid ON videos(video_id);
CREATE INDEX IF NOT EXISTS idx_videos_run ON videos(run_id);
"""


class DB:
    """Kết nối SQLite dùng chung cho luồng giao diện và luồng phân tích (có khóa)."""

    def __init__(self, path=DB_PATH):
        self.path = path
        self.lock = threading.RLock()
        self.conn = sqlite3.connect(path, check_same_thread=False)
        self.conn.row_factory = sqlite3.Row
        with self.lock:
            self.conn.executescript(SCHEMA)
            cols = {r[1] for r in self.conn.execute("PRAGMA table_info(videos)")}
            if "ai_input" not in cols:   # nội dung đã đọc (chờ dán kết quả AI) – thêm cho file cũ
                self.conn.execute("ALTER TABLE videos ADD COLUMN ai_input TEXT")
            self.conn.commit()

    # ---------------------------------------------------------------- runs
    def new_run(self, account, target, keywords, settings):
        with self.lock:
            cur = self.conn.execute(
                "INSERT INTO runs(started_at, account, target, keywords, settings) VALUES (?,?,?,?,?)",
                (_now(), account, target, ", ".join(keywords or []), json.dumps(settings or {}, ensure_ascii=False)))
            self.conn.commit()
            return cur.lastrowid

    def run_settings(self, run_id):
        if not run_id:
            return {}
        with self.lock:
            r = self.conn.execute("SELECT account, settings FROM runs WHERE id=?", (run_id,)).fetchone()
        if not r:
            return {}
        try:
            d = json.loads(r["settings"] or "{}")
        except Exception:
            d = {}
        d.setdefault("account", r["account"])
        return d

    def finish_run(self, run_id, note=""):
        with self.lock:
            self.conn.execute("UPDATE runs SET finished_at=?, note=? WHERE id=?", (_now(), note, run_id))
            self.conn.commit()

    # -------------------------------------------------------------- videos
    def add_video(self, run_id, v):
        with self.lock:
            cur = self.conn.execute(
                "INSERT INTO videos(run_id, video_id, url, title, channel, channel_url, duration, views, lang, "
                "relevance, status, found_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                (run_id, v.get("video_id"), v.get("url"), v.get("title"), v.get("channel"), v.get("channel_url"),
                 v.get("duration") or 0, v.get("views") or 0, v.get("lang", ""), v.get("relevance", 0),
                 v.get("status", "Đã tìm thấy"), _now()))
            self.conn.commit()
            return cur.lastrowid

    def update(self, row_id, **fields):
        if not fields:
            return
        for k in ("reasons", "comments"):
            if k in fields and not isinstance(fields[k], str):
                fields[k] = json.dumps(fields[k], ensure_ascii=False)
        keys = ", ".join(f"{k}=?" for k in fields)
        with self.lock:
            self.conn.execute(f"UPDATE videos SET {keys} WHERE id=?", (*fields.values(), row_id))
            self.conn.commit()

    def get(self, row_id):
        with self.lock:
            r = self.conn.execute("SELECT * FROM videos WHERE id=?", (row_id,)).fetchone()
        return _row(r) if r else None

    def rows(self, run_id=None, limit=2000):
        with self.lock:
            if run_id:
                rs = self.conn.execute("SELECT * FROM videos WHERE run_id=? ORDER BY id", (run_id,)).fetchall()
            else:
                rs = self.conn.execute("SELECT * FROM videos ORDER BY id DESC LIMIT ?", (limit,)).fetchall()
        return [_row(r) for r in rs]

    def seen_ids(self, days=30):
        """Mã video đã phân tích trong `days` ngày gần đây (để không lặp lại)."""
        with self.lock:
            rs = self.conn.execute(
                "SELECT DISTINCT video_id FROM videos WHERE analyzed_at IS NOT NULL "
                "AND analyzed_at >= datetime('now', 'localtime', ?)", (f"-{int(days)} days",)).fetchall()
        return {r[0] for r in rs if r[0]}

    def recent_comments(self, limit=300):
        """Bình luận đã tạo gần đây (để tránh viết lại câu giống nhau)."""
        out = []
        with self.lock:
            rs = self.conn.execute("SELECT comments FROM videos WHERE comments IS NOT NULL "
                                   "ORDER BY id DESC LIMIT ?", (limit,)).fetchall()
        for r in rs:
            try:
                out += [c for c in json.loads(r[0] or "[]") if isinstance(c, str)]
            except Exception:
                pass
        return out

    def pending_manual(self, run_id=None):
        """Video đã đọc nội dung, đang chờ dán kết quả AI (cách dán tay)."""
        with self.lock:
            q = "SELECT * FROM videos WHERE status='Chờ AI (dán tay)' AND ai_input IS NOT NULL"
            rs = self.conn.execute(q + (" AND run_id=?" if run_id else "") + " ORDER BY id",
                                   (run_id,) if run_id else ()).fetchall()
        return [_row(r) for r in rs]

    def clear(self):
        with self.lock:
            self.conn.execute("DELETE FROM videos")
            self.conn.execute("DELETE FROM runs")
            self.conn.commit()

    # -------------------------------------------------------------- xuất
    def export(self, path, rows):
        """Xuất báo cáo .csv hoặc .xlsx (theo đuôi file)."""
        head = ["STT", "Video", "Link", "Kênh", "Chủ đề", "Điểm", "Phân loại", "Đề xuất", "Trạng thái",
                "Tương tác đã làm", "Lý do", "Căn cứ Dislike", "Bình luận 1", "Bình luận 2", "Bình luận 3",
                "Ngôn ngữ bình luận", "Khoảng phân tích", "Nguồn nội dung", "Độ liên quan", "Thời điểm"]
        data = []
        for i, r in enumerate(rows, 1):
            cm = (r.get("comments") or []) + ["", "", ""]
            data.append([i, r.get("title"), r.get("url"), r.get("channel"), r.get("topic"), r.get("score"),
                         r.get("category"), r.get("recommendation"), r.get("status"), r.get("action") or "",
                         " • ".join(r.get("reasons") or []), r.get("dislike_basis") or "", cm[0], cm[1], cm[2],
                         r.get("comment_lang") or "", _window(r), r.get("transcript_source") or "",
                         r.get("relevance"), r.get("analyzed_at") or r.get("found_at")])
        if path.lower().endswith(".xlsx"):
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            wb = Workbook()
            ws = wb.active
            ws.title = "Phân tích video"
            ws.append(head)
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="2F3542")
            for row in data:
                ws.append(row)
            widths = [6, 48, 44, 24, 22, 8, 12, 10, 16, 16, 60, 30, 50, 50, 50, 14, 14, 14, 10, 20]
            for i, w in enumerate(widths, 1):
                ws.column_dimensions[ws.cell(1, i).column_letter].width = w
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.freeze_panes = "A2"
            wb.save(path)
        else:
            with open(path, "w", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f)
                w.writerow(head)
                w.writerows(data)
        return path


def _window(r):
    a, b = r.get("window_from"), r.get("window_to")
    if a is None or b is None:
        return ""
    return f"{a // 60}:{a % 60:02d}–{b // 60}:{b % 60:02d}"


def _row(r):
    d = dict(r)
    for k in ("reasons", "comments"):
        try:
            d[k] = json.loads(d.get(k) or "[]")
        except Exception:
            d[k] = []
    return d


def _now():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")
