"""
Trang "Kết quả" dễ đọc trong file ket_qua_keo_kenh.xlsx.

File có 2 trang:
  · "Kết quả"      – trang bạn xem: mỗi thao tác 1 dòng, chữ ngắn, mới nhất ở trên cùng, tô màu theo kết quả.
  · "Dữ liệu gốc"  – nhật ký đầy đủ tool dùng để đối soát / gỡ (giữ nguyên, không cần đọc).
Trang "Kết quả" được dựng lại từ "Dữ liệu gốc" mỗi lần tool ghi – chỉ rút gọn chữ, không suy đoán thêm gì.
"""
import re
from datetime import datetime

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

DATA_NAME = "Dữ liệu gốc"
OLD_NAMES = ("Nhat ky keo kenh",)
VIEW_NAME = "Kết quả"

FONT = "Arial"
INK, MUTED = "1F2A37", "6B7785"
HEAD_BG = "1F3B57"
ZEBRA = PatternFill("solid", start_color="F6F8FB")
LINE = Side(style="thin", color="E2E7ED")
DAY_LINE = Side(style="medium", color="9AA8B8")

RESULT_STYLE = {   # nhãn -> (màu chữ, màu nền)
    "✅ Xong": ("1E7B34", "E3F4E7"),
    "⏭ Có sẵn": ("3B5B85", "E6EEF8"),
    "⚠️ Xem lại": ("8A5A00", "FFF1CC"),
    "❌ Lỗi": ("B42318", "FDE4E2"),
}

TASK_SHORT = {
    "Màn hình kết thúc": "Màn hình kết thúc",
    "Thẻ (cards)": "Thẻ",
    "Bình luận ghim": "Bình luận ghim",
    "Thẻ Trang chủ": "Trang chủ kênh",
    "Hồ sơ: đường liên kết": "Link trên hồ sơ",
    "Bài đăng video": "Bài đăng",
    "Tính năng cộng tác": "Cộng tác",
    "Đọc link": "Đọc link",
}

COLS = [("Lúc", 13), ("Kênh mạnh", 22), ("Kéo cho kênh", 22), ("Video", 15), ("Việc", 17),
        ("Kết quả", 12), ("Ghi chú", 48), ("Gỡ", 11)]
HEAD_ROW = 4          # dòng tiêu đề cột (dòng 1–2: tên bảng + tổng kết)
NOTE_MAX = 60


# ------------------------------------------------------------------ tìm trang dữ liệu
def log_sheet(wb):
    """Trang nhật ký đầy đủ (tool đọc / ghi). Hỗ trợ cả file cũ chỉ có 1 trang."""
    for n in (DATA_NAME,) + OLD_NAMES:
        if n in wb.sheetnames:
            return wb[n]
    for ws in wb.worksheets:
        if ws.title != VIEW_NAME:
            return ws
    return wb.active


# ------------------------------------------------------------------ rút gọn chữ
def _cut(s, n=NOTE_MAX):
    s = re.sub(r"\s+", " ", str(s or "")).strip(" .|–-")
    return s if len(s) <= n else s[:n - 1].rstrip(" ,.–-") + "…"


def _name(s):
    """'Tên kênh (UC...)' -> 'Tên kênh'."""
    s = re.sub(r"\s*\((UC[\w-]{22})?\)\s*$", "", str(s or "")).strip()
    return _cut(s, 24) if s else ""


def _video(target):
    """(chữ hiện, link) – mã video cho link video, 'Trang kênh' cho link kênh."""
    t = str(target or "")
    m = re.search(r"(?:v=|youtu\.be/|/shorts/|/video/)([\w-]{11})", t)
    if m:
        return m.group(1), f"https://www.youtube.com/watch?v={m.group(1)}"
    if "/posts" in t:
        return "Tab Bài đăng", t
    if t.startswith("http"):
        return "Trang kênh", t
    return _cut(t, 15), ""


def _main_error(err):
    """Bỏ các phần 'GỠ: …' (ghi chú lúc gỡ) khỏi cột Lỗi."""
    parts = [p.strip() for p in str(err or "").split(" | ")]
    keep = [p for p in parts if p and not p.startswith("GỠ:")]
    return " | ".join(keep)


def _already(detail):
    d = str(detail or "")
    return bool(re.search(r"đã có sẵn|đã đúng sẵn", d, re.I)) and "Đã làm:" not in d and "Đã thêm" not in d \
        and "Đã xóa" not in d


def _items(label, text):
    m = re.search(label + r":\s*([^.]+)", text)
    return [x.strip() for x in m.group(1).split(",") if x.strip()] if m else []


_REASONS = [   # (mẫu trong lời báo lỗi gốc, câu ngắn)
    (r"Timeout|TimeoutError", "Trang tải chậm, quá thời gian chờ – chạy lại"),
    (r"trình chỉnh sửa chưa hiện video đó", "Chọn video chưa được – chạy lại"),
    (r"Menu ⋮ không có mục 'Xóa'", "Không xóa được bình luận cũ"),
    (r"chưa xác nhận được bình luận mới đang được ghim", "Đã đăng, chưa chắc đã ghim – mở video xem"),
    (r"Rất tiếc", "Studio báo 'Rất tiếc' – video lỗi / đã xóa"),
    (r"đang bị khóa", "Mục bị khóa (video trẻ em / Shorts / quá ngắn)"),
    (r"chỉ có (\d+) link video", "Kênh cần kéo chỉ đủ {0} video"),
    (r"hộp chọn video không mở", "Hộp chọn video không mở ra"),
    (r"Menu 'Thêm phần kênh' không có mục", "YouTube không cho thêm mục này"),
    (r"Cộng tác chưa hiện tên kênh", "Đã mời, chưa thấy tên kênh – mở video xem"),
    (r"chưa bấm được Chấp nhận", "Đã mời – cần bấm Chấp nhận bên kênh kéo"),
    (r"không thấy nút 'Mời cộng sự", "Không thấy nút Mời cộng sự"),
    (r"khung bài đăng chưa hiện video", "Chưa gắn được video vào bài – đã hủy"),
    (r"không thấy ô thời gian", "Không đặt được mốc thời gian của thẻ"),
    (r"Không mở được '([^']+)'", "Không mở được trang {0}"),
    (r"Không tìm thấy nút/mục", "Không thấy nút trên trang (YouTube đổi giao diện?)"),
    (r"không vào được kênh|không có kênh này", "Sai tài khoản / không vào được kênh"),
]


def _reason(text):
    t = str(text or "")
    for pat, short in _REASONS:
        m = re.search(pat, t, re.I)
        if m:
            return short.format(*m.groups())
    t = re.sub(r"^\[[^\]]*\]\s*", "", t)
    return _cut(re.split(r"(?<=[.!?])\s", t)[0])


_HOME_PARTS = [("giới thiệu", "Video giới thiệu"), ("nổi bật cho người", "Video nổi bật"),
               ("nổi bật dành cho", "Video nổi bật"), ("Danh sách phát", "Danh sách phát"),
               ("Kênh nổi bật", "Kênh nổi bật")]


def _part(s):
    for k, v in _HOME_PARTS:
        if k.lower() in s.lower():
            return v
    return _cut(s, 24)


def _home_note(text, ok):
    """'Đã làm: A, B. Đã có sẵn: C. X: lý do | …' -> 'Làm 2 mục · 1 có sẵn · Thiếu: Video giới thiệu'."""
    done, have = _items("Đã làm", text), _items("Đã có sẵn", text)
    bits = [f"Làm {len(done)} mục"] if done else []
    if have:
        bits.append(f"{len(have)} có sẵn")
    if not ok:
        rest = re.sub(r"(Đã làm|Đã có sẵn|Không làm[^:]*):[^.]*\.?", "", text.split(" | ")[0]).strip(" .")
        if ":" in rest:
            bits.append("Thiếu: " + _part(rest.split(":")[0]))
    return " · ".join(bits)


def _note(task, detail, err, status, kind):
    detail, main_err = str(detail or ""), _main_error(err)
    if kind == "⏭ Có sẵn":
        return "Đã có từ trước – giữ nguyên"
    if kind in ("❌ Lỗi", "⚠️ Xem lại"):
        if task == "Thẻ Trang chủ":
            n = _home_note(main_err or detail, False)
            return _cut(n if "Thiếu" in n else (n + " · Mở trang chủ kênh xem lại" if n else "Mở trang chủ kênh xem lại"))
        src = main_err.split(" — ", 1)[1] if " — " in main_err else main_err
        if task == "Thẻ (cards)" and not src.strip() or re.match(r"^\d+:\d\d ", src):
            return "Mở Studio kiểm tra lại thẻ"
        return _reason(src) if src else "Mở Studio kiểm tra lại"
    # ✅ Xong
    ids = list(dict.fromkeys(re.findall(r"\(([\w-]{11})\)", detail)))
    if task == "Thẻ (cards)":
        marks = re.findall(r"(?:^|\| )(\d+:\d\d) ", detail)
        pl = " + danh sách phát" if re.search(r"\((?:PL|OL|UU)[\w-]{10,}\)", detail) else ""
        return f"Gắn {len(marks) or len(ids)} thẻ" + (f" ({' · '.join(marks)})" if marks else "") + pl
    if task == "Màn hình kết thúc":
        pl = " + danh sách phát" if "Danh sách phát:" in detail else ""
        return (f"Gắn {len(ids)} video" if ids else "Đã gắn video") + pl
    if task == "Bình luận ghim":
        return "Đăng + ghim bình luận"
    if task == "Thẻ Trang chủ":
        return _home_note(detail, True) or "Đã cập nhật"
    if task == "Hồ sơ: đường liên kết":
        m = re.search(r"Đã thêm đường liên kết:\s*(.+?)\s+–\s+http", detail)
        pre = "Thay link cũ → " if "Đã xóa đường liên kết cũ" in detail else "Thêm link: "
        return _cut(pre + m.group(1)) if m else ("Thay link cũ" if pre.startswith("Thay") else "Đã thêm link")
    if task == "Bài đăng video":
        return "Đã đăng bài" + (f" (video {ids[0]})" if ids else "")
    if task == "Tính năng cộng tác":
        return "Đã mời cộng tác"
    return _cut(detail) if detail else ""


def _kind(status, detail):
    s = str(status or "")
    if "LỖI" in s:
        return "❌ Lỗi"
    if "KIỂM" in s:
        return "⚠️ Xem lại"
    if "DONE" in s:
        return "⏭ Có sẵn" if _already(detail) else "✅ Xong"
    return ""


def _removal(v, status):
    v = str(v or "")
    if "Đã gỡ" in v:
        return "✅ Đã gỡ"
    if "Lỗi" in v:
        return "❌ Lỗi gỡ"
    if "thủ công" in v:
        return "⚠️ Gỡ tay"
    if "Bỏ qua" in v:
        return "⏭ Bỏ qua"
    return ""   # chưa gỡ -> để trống cho thoáng


def _when(t):
    try:
        d = datetime.strptime(str(t), "%Y-%m-%d %H:%M:%S")
        return d, d.strftime("%d/%m  %H:%M")
    except Exception:
        return None, str(t or "")


# ------------------------------------------------------------------ dựng trang
def build_view(wb):
    """Dựng lại trang 'Kết quả' (đứng đầu, mở file là thấy ngay). Trả về trang đó."""
    data = log_sheet(wb)
    if data.title in OLD_NAMES:
        data.title = DATA_NAME
    head = [c.value for c in data[1]]
    col = {h: i for i, h in enumerate(head) if h}
    g = lambda row, k: row[col[k]] if k in col and col[k] < len(row) else None

    recs = []
    for row in data.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        status, detail, task = g(row, "Trạng thái"), g(row, "Chi tiết"), str(g(row, "Hạng mục") or "")
        kind = _kind(status, detail)
        d, when = _when(g(row, "Thời gian"))
        vid, link = _video(g(row, "Link video / Kênh"))
        strong = _name(g(row, "Kênh mạnh (dùng để kéo)")) or _name(g(row, "Kênh"))
        weak = _name(g(row, "Kênh cần kéo"))
        recs.append(dict(d=d, when=when, strong=strong, weak=weak if weak and weak != strong else "Chính kênh", vid=vid, link=link,
                         task=TASK_SHORT.get(task, _cut(task, 18)), kind=kind,
                         note=_note(task, detail, g(row, "Lỗi"), status, kind),
                         gone=_removal(g(row, "Tình trạng gỡ"), status)))
    recs.reverse()   # mới nhất ở trên

    if VIEW_NAME in wb.sheetnames:
        del wb[VIEW_NAME]
    ws = wb.create_sheet(VIEW_NAME, 0)
    wb.active = 0
    for s in wb.worksheets:
        s.sheet_view.tabSelected = s is ws
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "2F6FE0"
    data.sheet_properties.tabColor = "A0AAB5"

    # dòng 1–2: tên bảng + tổng kết
    cnt = {k: sum(1 for r in recs if r["kind"] == k) for k in RESULT_STYLE}
    ws["A1"] = "KẾT QUẢ KÉO KÊNH"
    ws["A1"].font = Font(name=FONT, size=15, bold=True, color=INK)
    ws["A2"] = (f"Cập nhật {datetime.now().strftime('%d/%m/%Y %H:%M')}     "
                f"✅ Xong {cnt['✅ Xong']}     ⏭ Có sẵn {cnt['⏭ Có sẵn']}     "
                f"⚠️ Xem lại {cnt['⚠️ Xem lại']}     ❌ Lỗi {cnt['❌ Lỗi']}")
    ws["A2"].font = Font(name=FONT, size=10, color=MUTED)
    ws["A3"] = "Mới nhất ở trên cùng · Bấm mã video để mở · Chi tiết đầy đủ ở trang “Dữ liệu gốc”"
    ws["A3"].font = Font(name=FONT, size=9, italic=True, color=MUTED)
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 18

    for i, (h, w) in enumerate(COLS, 1):
        c = ws.cell(HEAD_ROW, i, h)
        c.font = Font(name=FONT, size=10.5, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color=HEAD_BG)
        c.alignment = Alignment(horizontal="center" if h in ("Lúc", "Kết quả", "Gỡ", "Video") else "left",
                                vertical="center", indent=0 if h in ("Lúc", "Kết quả", "Gỡ", "Video") else 1)
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[HEAD_ROW].height = 26

    base = Font(name=FONT, size=10, color=INK)
    muted = Font(name=FONT, size=10, color=MUTED)
    prev_day = None
    for n, r in enumerate(recs):
        rr = HEAD_ROW + 1 + n
        vals = [r["when"], r["strong"], r["weak"], r["vid"], r["task"], r["kind"], r["note"], r["gone"]]
        day = r["d"].date() if r["d"] else None
        top = DAY_LINE if (prev_day is not None and day != prev_day) else None
        prev_day = day
        for i, v in enumerate(vals, 1):
            c = ws.cell(rr, i, v)
            c.font = muted if i in (1, 8) or v == "Chính kênh" else base
            c.alignment = Alignment(horizontal="center" if i in (1, 4, 6, 8) else "left", vertical="center",
                                    indent=0 if i in (1, 4, 6, 8) else 1)
            c.border = Border(bottom=LINE, top=top) if top else Border(bottom=LINE)
            if n % 2:
                c.fill = ZEBRA
        if r["link"]:
            c = ws.cell(rr, 4)
            c.hyperlink = r["link"]
            c.font = Font(name=FONT, size=10, color="1F5FBF", underline="single")
        st = RESULT_STYLE.get(r["kind"])
        if st:
            c = ws.cell(rr, 6)
            c.font = Font(name=FONT, size=10, bold=True, color=st[0])
            c.fill = PatternFill("solid", start_color=st[1])
        ws.row_dimensions[rr].height = 21

    last = HEAD_ROW + max(len(recs), 1)
    ws.freeze_panes = ws.cell(HEAD_ROW + 1, 1)
    ws.auto_filter.ref = f"A{HEAD_ROW}:{ws.cell(HEAD_ROW, len(COLS)).column_letter}{last}"
    ws.sheet_view.zoomScale = 110
    return ws


def tidy_data_sheet(ws):
    """Trang 'Dữ liệu gốc': không xuống dòng, ẩn các cột kỹ thuật (vẫn còn đó, bỏ ẩn được)."""
    head = [c.value for c in ws[1]]
    for name in ("Kênh", "File kênh", "Dữ liệu gắn (JSON)"):
        if name in head:
            ws.column_dimensions[ws.cell(1, head.index(name) + 1).column_letter].hidden = True
    ws.freeze_panes = "A2"


def upgrade_file(path):
    """Chuyển file kết quả cũ sang dạng mới (thêm trang 'Kết quả') ngay khi mở tool / bấm mở file.
    'ok' = đã chuyển, 'same' = đã đúng dạng mới, 'busy' = file đang mở trong Excel/WPS, 'none' = chưa có file."""
    import os
    if not os.path.exists(path):
        return "none"
    from openpyxl import load_workbook
    try:
        wb = load_workbook(path, read_only=True)
        names = list(wb.sheetnames)
        n_data = log_sheet(wb).max_row or 0
        n_view = wb[VIEW_NAME].max_row if VIEW_NAME in names else -1
        wb.close()
        if VIEW_NAME in names and names[0] == VIEW_NAME and n_view == HEAD_ROW + max(n_data - 1, 1):
            return "same"
        wb = load_workbook(path)
        tidy_data_sheet(log_sheet(wb))
        build_view(wb)
        wb.save(path)
        return "ok"
    except PermissionError:
        return "busy"


# ------------------------------------------------------------------ khi file Excel đang mở trong Excel / WPS
def pending_path(path):
    """File phụ tool ghi tạm khi file chính đang bị Excel / WPS giữ (không ghi đè được)."""
    import os
    root, ext = os.path.splitext(path)
    return root + "_CHO_GHI" + ext


def latest(path):
    """Bản nhật ký MỚI NHẤT để đọc: file phụ (nếu có và mới hơn) hoặc file chính."""
    import os
    p = pending_path(path)
    if os.path.exists(p) and (not os.path.exists(path) or os.path.getmtime(p) >= os.path.getmtime(path)):
        return p
    return path


def merge_pending(path):
    """Gộp file phụ vào file chính. True = xong (hoặc không có gì để gộp), False = file chính vẫn đang mở."""
    import os
    p = pending_path(path)
    if not os.path.exists(p):
        return True
    try:
        if os.path.exists(path) and os.path.getmtime(p) < os.path.getmtime(path):
            os.remove(p)          # file phụ cũ hơn file chính -> bỏ
        else:
            os.replace(p, path)
        return True
    except PermissionError:
        return False
    except OSError:
        return False


def drop_pending(path):
    import os
    try:
        os.remove(pending_path(path))
    except OSError:
        pass


def is_locked(path):
    """True nếu file đang bị Excel / WPS mở giữ (tool không ghi được)."""
    import os
    if not os.path.exists(path):
        return False
    try:
        with open(path, "r+b"):
            return False
    except PermissionError:
        return True
    except OSError:
        return False
